"""집계 기간 전체의 통계 엑셀 워크북을 만든다. .venv 로 실행할 것.

시트: 요약 / 매체×유형 / 월별추이 / 주별추이 / 성과 / 상위콘텐츠 / 활동목록
수집 결과(snapshots/parts)를 그대로 읽으므로, 최신 수치가 필요하면 수집기를 먼저 돌린다.
"""
import os
import statistics
import sys
from collections import Counter
from datetime import datetime, timedelta

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from analytics.common import OUT_DIR, log  # noqa: E402
from analytics.build_report import MED, TYP, merge  # noqa: E402

from openpyxl import Workbook                                    # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter                     # noqa: E402

SINCE = "2026-03-01"
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
HEAD_FILL = PatternFill("solid", fgColor="3F4448")
SUB_FILL = PatternFill("solid", fgColor="EDEDEA")
TITLE = Font(bold=True, size=12)
BOLD = Font(bold=True)
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))


def head_row(ws, row, values, width=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row, i, v)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
    return row + 1


def section(ws, row, text):
    ws.cell(row, 1, text).font = TITLE
    return row + 1


def widths(ws, spec):
    for i, w in enumerate(spec, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build(rows, channel, profile, path):
    rows = [r for r in rows if r["date"] >= SINCE]
    made = [r for r in rows if r["type"] != "게시글"]
    months = sorted({r["date"][:7] for r in rows})
    last = max(r["date"] for r in rows)
    wb = Workbook()

    # ---------------------------------------------------------------- 요약
    ws = wb.active
    ws.title = "요약"
    r = 1
    ws.cell(r, 1, "세명대학교 재난안전학과 콘텐츠 활동 통계").font = Font(bold=True, size=14)
    r += 1
    ws.cell(r, 1, f"집계 기간 {SINCE} ~ {last} · 4개 매체 · 유형 5종")
    r += 2
    r = section(ws, r, "총괄")
    r = head_row(ws, r, ["구분", "건수", "비중"])
    total = len(rows)
    for label, n in [("전체", total), ("제작 콘텐츠(영상·숏츠·이미지·카드뉴스)", len(made)),
                     ("게시글(공지·블로그)", total - len(made))]:
        ws.cell(r, 1, label)
        ws.cell(r, 2, n)
        ws.cell(r, 3, n / total).number_format = "0.0%"
        r += 1
    r += 1
    r = section(ws, r, "매체별")
    r = head_row(ws, r, ["매체", "건수", "비중", "제작 콘텐츠"])
    for m in MED:
        n = sum(1 for x in rows if x["media"] == m)
        ws.cell(r, 1, m)
        ws.cell(r, 2, n)
        ws.cell(r, 3, n / total).number_format = "0.0%"
        ws.cell(r, 4, sum(1 for x in made if x["media"] == m))
        r += 1
    r += 1
    r = section(ws, r, "유형별")
    r = head_row(ws, r, ["유형", "건수", "비중"])
    tc = Counter(x["type"] for x in rows)
    for t in TYP:
        ws.cell(r, 1, t)
        ws.cell(r, 2, tc.get(t, 0))
        ws.cell(r, 3, tc.get(t, 0) / total).number_format = "0.0%"
        r += 1
    r += 1
    r = section(ws, r, "채널 규모")
    r = head_row(ws, r, ["지표", "값"])
    for k, v in [("유튜브 구독자", channel.get("subscribers")),
                 ("유튜브 채널 누적 조회", channel.get("total_views")),
                 ("인스타그램 팔로워", profile.get("followers")),
                 ("인스타그램 총 게시물", profile.get("posts"))]:
        ws.cell(r, 1, k)
        ws.cell(r, 2, v if v is not None else "—").number_format = "#,##0"
        r += 1
    widths(ws, [38, 14, 10, 14])

    # ------------------------------------------------------------ 매체×유형
    ws = wb.create_sheet("매체×유형")
    r = head_row(ws, 1, ["매체"] + TYP + ["합계"])
    for m in MED:
        rs = [x for x in rows if x["media"] == m]
        c = Counter(x["type"] for x in rs)
        ws.cell(r, 1, m)
        for i, t in enumerate(TYP, 2):
            ws.cell(r, i, c.get(t, 0))
        ws.cell(r, len(TYP) + 2, len(rs)).font = BOLD
        r += 1
    ws.cell(r, 1, "합계").font = BOLD
    for i, t in enumerate(TYP, 2):
        ws.cell(r, i, tc.get(t, 0)).font = BOLD
    ws.cell(r, len(TYP) + 2, total).font = BOLD
    widths(ws, [16] + [10] * (len(TYP) + 1))
    ws.freeze_panes = "B2"

    # ------------------------------------------------------------- 월별 추이
    ws = wb.create_sheet("월별추이")
    r = section(ws, 1, "매체별")
    r = head_row(ws, r, ["월"] + MED + ["합계", "제작 콘텐츠"])
    for mo in months:
        ws.cell(r, 1, mo)
        v = [sum(1 for x in rows if x["date"][:7] == mo and x["media"] == m) for m in MED]
        for i, n in enumerate(v, 2):
            ws.cell(r, i, n)
        ws.cell(r, len(MED) + 2, sum(v)).font = BOLD
        ws.cell(r, len(MED) + 3, sum(1 for x in made if x["date"][:7] == mo))
        r += 1
    r += 1
    r = section(ws, r, "유형별")
    r = head_row(ws, r, ["월"] + TYP + ["합계"])
    for mo in months:
        ws.cell(r, 1, mo)
        c = Counter(x["type"] for x in rows if x["date"][:7] == mo)
        for i, t in enumerate(TYP, 2):
            ws.cell(r, i, c.get(t, 0))
        ws.cell(r, len(TYP) + 2, sum(c.values())).font = BOLD
        r += 1
    widths(ws, [12] + [14] * 6)

    # ------------------------------------------------------------- 주별 추이
    ws = wb.create_sheet("주별추이")
    r = head_row(ws, 1, ["주 시작(월)", "유튜브", "인스타그램", "네이버 블로그", "학과 홈페이지",
                         "합계", "제작 콘텐츠"])
    d = datetime.strptime(SINCE, "%Y-%m-%d")
    d -= timedelta(days=d.weekday())                 # 그 주 월요일부터
    end = datetime.strptime(last, "%Y-%m-%d")
    while d <= end:
        s, e = d.strftime("%Y-%m-%d"), (d + timedelta(days=7)).strftime("%Y-%m-%d")
        wk = [x for x in rows if s <= x["date"] < e]
        ws.cell(r, 1, s)
        for i, m in enumerate(MED, 2):
            ws.cell(r, i, sum(1 for x in wk if x["media"] == m))
        ws.cell(r, 6, len(wk)).font = BOLD
        ws.cell(r, 7, sum(1 for x in wk if x["type"] != "게시글"))
        r += 1
        d += timedelta(days=7)
    widths(ws, [14, 12, 14, 16, 16, 10, 14])
    ws.freeze_panes = "B2"

    # ----------------------------------------------------------------- 성과
    ws = wb.create_sheet("성과")
    yt = [x for x in rows if x["media"] == "유튜브"]
    ig = [x for x in rows if x["media"] == "인스타그램"]
    v = [x.get("views", 0) for x in yt]
    r = section(ws, 1, "유튜브 조회수")
    r = head_row(ws, r, ["구분", "편수", "누적", "평균", "중앙값", "최고"])
    for label, sel in [("전체", yt), ("숏츠", [x for x in yt if x["type"] == "숏츠"]),
                       ("영상(60초 초과)", [x for x in yt if x["type"] == "영상"])]:
        vv = [x.get("views", 0) for x in sel]
        ws.cell(r, 1, label)
        ws.cell(r, 2, len(sel))
        for i, n in enumerate([sum(vv), round(statistics.mean(vv)) if vv else 0,
                               int(statistics.median(vv)) if vv else 0, max(vv) if vv else 0], 3):
            ws.cell(r, i, n).number_format = "#,##0"
        r += 1
    r += 1
    r = section(ws, r, "유튜브 월별 조회")
    r = head_row(ws, r, ["월", "편수", "누적 조회", "편당 평균"])
    for mo in months:
        sel = [x for x in yt if x["date"][:7] == mo]
        if not sel:
            continue
        vv = [x.get("views", 0) for x in sel]
        ws.cell(r, 1, mo)
        ws.cell(r, 2, len(sel))
        ws.cell(r, 3, sum(vv)).number_format = "#,##0"
        ws.cell(r, 4, round(statistics.mean(vv))).number_format = "#,##0"
        r += 1
    r += 1
    r = section(ws, r, "인스타그램 반응 (유형별)")
    r = head_row(ws, r, ["유형", "건수", "좋아요 평균", "좋아요 최대", "댓글 합계", "조회 평균"])
    for t in ["이미지", "카드뉴스", "숏츠", "영상"]:
        sel = [x for x in ig if x["type"] == t]
        if not sel:
            continue
        lk = [x.get("likes", 0) for x in sel]
        vw = [x.get("views", 0) for x in sel if x.get("views")]
        ws.cell(r, 1, t)
        ws.cell(r, 2, len(sel))
        ws.cell(r, 3, round(statistics.mean(lk), 1))
        ws.cell(r, 4, max(lk))
        ws.cell(r, 5, sum(x.get("comments", 0) for x in sel))
        ws.cell(r, 6, round(statistics.mean(vw)) if vw else 0).number_format = "#,##0"
        r += 1
    r += 1
    r = section(ws, r, "점검")
    r = head_row(ws, r, ["항목", "값"])
    empty = [x for x in ig if x.get("caption_empty")]
    ws.cell(r, 1, "인스타 캡션 누락")
    ws.cell(r, 2, len(empty))
    r += 1
    cap = [x.get("views", 0) for x in ig if x["type"] == "숏츠" and not x.get("caption_empty") and x.get("views")]
    nocap = [x.get("views", 0) for x in ig if x["type"] == "숏츠" and x.get("caption_empty") and x.get("views")]
    ws.cell(r, 1, "릴스 조회 평균 — 캡션 있음")
    ws.cell(r, 2, round(statistics.mean(cap)) if cap else 0)
    r += 1
    ws.cell(r, 1, "릴스 조회 평균 — 캡션 없음")
    ws.cell(r, 2, round(statistics.mean(nocap)) if nocap else 0)
    widths(ws, [28, 12, 14, 14, 14, 14])

    # ------------------------------------------------------------ 상위 콘텐츠
    ws = wb.create_sheet("상위콘텐츠")
    r = section(ws, 1, "유튜브 조회 상위 15")
    r = head_row(ws, r, ["순위", "조회수", "게시일", "유형", "제목", "링크"])
    for i, x in enumerate(sorted(yt, key=lambda a: -a.get("views", 0))[:15], 1):
        ws.cell(r, 1, i)
        ws.cell(r, 2, x.get("views", 0)).number_format = "#,##0"
        ws.cell(r, 3, x["date"])
        ws.cell(r, 4, x["type"])
        ws.cell(r, 5, x["title"])
        ws.cell(r, 6, x["url"])
        r += 1
    r += 1
    r = section(ws, r, "인스타그램 좋아요 상위 15")
    r = head_row(ws, r, ["순위", "좋아요", "게시일", "유형", "활동명", "링크"])
    for i, x in enumerate(sorted(ig, key=lambda a: -a.get("likes", 0))[:15], 1):
        ws.cell(r, 1, i)
        ws.cell(r, 2, x.get("likes", 0))
        ws.cell(r, 3, x["date"])
        ws.cell(r, 4, x["type"])
        ws.cell(r, 5, x["title"] or "(제목 없음)")
        ws.cell(r, 6, x["url"])
        r += 1
    widths(ws, [8, 12, 12, 12, 60, 46])

    # -------------------------------------------------------------- 활동목록
    ws = wb.create_sheet("활동목록")
    head_row(ws, 1, ["매체명", "유형", "활동명", "게시일", "조회수", "좋아요", "링크"])
    r = 2
    for x in sorted(rows, key=lambda a: a["date"], reverse=True):
        ws.cell(r, 1, x["media"])
        ws.cell(r, 2, x["type"])
        ws.cell(r, 3, x["title"] or "(제목 없음)")
        ws.cell(r, 4, x["date"])
        ws.cell(r, 5, x.get("views") or None).number_format = "#,##0"
        ws.cell(r, 6, x.get("likes") or None)
        ws.cell(r, 7, x["url"])
        for c in range(1, 8):
            ws.cell(r, c).border = THIN
        r += 1
    widths(ws, [14, 10, 62, 12, 10, 9, 46])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{ws.max_row}"

    wb.save(path)
    return {"total": total, "made": len(made), "last": last}


if __name__ == "__main__":
    rows, channel, profile, warn = merge()
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        OUT_DIR, f"재난안전학과_콘텐츠활동_통계_{datetime.now():%Y%m%d}.xlsx")
    info = build(rows, channel, profile, out)
    log(f"통계 엑셀 생성: {out} (전체 {info['total']}건, 제작 {info['made']}건, ~{info['last']})")
    print(out)
